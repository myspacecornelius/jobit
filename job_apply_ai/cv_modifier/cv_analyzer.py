"""CV analyzer and modifier.

Extracts skills from job descriptions and rewrites the skills section of a
.docx CV template. Config comes from `job_apply_ai.config` — notably the
spaCy model name and whether to auto-download it.
"""

from __future__ import annotations

import os
import re
import subprocess
import sys
from datetime import datetime
from typing import Dict, List, Optional, Tuple

import pandas as pd
import spacy
from docx import Document

from job_apply_ai.config import get_config
from job_apply_ai.utils.helpers import ensure_directory_exists, sanitize_filename
from job_apply_ai.utils.logging import get_logger

logger = get_logger(__name__)


class MissingSpacyModelError(RuntimeError):
    """Raised when the configured spaCy model is not installed."""


def _load_spacy_model(model_name: str, auto_download: bool):
    """Load the spaCy model, optionally downloading it if missing."""
    try:
        return spacy.load(model_name)
    except OSError as exc:
        if not auto_download:
            raise MissingSpacyModelError(
                f"spaCy model '{model_name}' is not installed. Run:\n"
                f"    python -m spacy download {model_name}\n"
                f"or set JOBIT_SPACY_AUTO_DOWNLOAD=true in your .env."
            ) from exc
        logger.info("Downloading spaCy model '%s'...", model_name)
        subprocess.run(
            [sys.executable, "-m", "spacy", "download", model_name],
            check=True,
        )
        return spacy.load(model_name)


SKILL_CATEGORIES: Dict[str, List[str]] = {
    "Programming Languages": [
        "python", "java", "javascript", "c++", "c#", "ruby", "php", "swift",
        "kotlin", "go", "rust", "typescript", "scala", "perl", "r", "matlab",
        "bash", "shell", "powershell", "sql", "html", "css", "dart",
    ],
    "Frameworks & Libraries": [
        "react", "angular", "vue", "django", "flask", "spring", "express",
        "node.js", "tensorflow", "pytorch", "scikit-learn", "pandas", "numpy",
        "bootstrap", "jquery", "laravel", "symfony", "rails", "asp.net",
        "flutter", "xamarin", ".net", "dotnet", "core", "entity framework",
    ],
    "Databases": [
        "mysql", "postgresql", "mongodb", "sqlite", "oracle", "sql server",
        "cassandra", "redis", "elasticsearch", "dynamodb", "mariadb", "neo4j",
        "firebase", "supabase", "cockroachdb", "couchdb", "cosmosdb",
    ],
    "Cloud & DevOps": [
        "aws", "azure", "gcp", "google cloud", "docker", "kubernetes", "jenkins",
        "terraform", "ansible", "chef", "puppet", "circleci", "travis", "github actions",
        "gitlab ci", "bitbucket pipelines", "heroku", "netlify", "vercel", "digitalocean",
        "linode", "cloudflare", "akamai", "fastly", "lambda", "ec2", "s3", "rds",
    ],
    "Tools & Platforms": [
        "git", "github", "gitlab", "bitbucket", "jira", "confluence", "trello",
        "slack", "notion", "figma", "sketch", "adobe xd", "photoshop", "illustrator",
        "visual studio", "vs code", "intellij", "pycharm", "eclipse", "android studio",
        "xcode", "postman", "insomnia", "swagger", "sentry", "datadog", "grafana",
    ],
    "Methodologies": [
        "agile", "scrum", "kanban", "waterfall", "lean", "tdd", "bdd", "ci/cd",
        "devops", "devsecops", "gitflow", "trunk-based development", "pair programming",
        "extreme programming", "safe", "prince2", "pmp", "itil", "togaf",
    ],
    "Soft Skills": [
        "communication", "teamwork", "leadership", "problem solving", "critical thinking",
        "time management", "adaptability", "creativity", "emotional intelligence",
        "conflict resolution", "negotiation", "presentation", "public speaking",
        "customer service", "mentoring", "coaching", "decision making",
    ],
    "Languages": [
        "english", "german", "french", "spanish", "italian", "portuguese", "dutch",
        "swedish", "norwegian", "danish", "finnish", "russian", "chinese", "japanese",
        "korean", "arabic", "hindi", "bengali", "urdu", "turkish", "polish", "ukrainian",
    ],
    "Business & Analytics": [
        "excel", "powerpoint", "word", "tableau", "power bi", "looker", "google analytics",
        "seo", "sem", "google ads", "facebook ads", "marketing", "sales", "crm", "erp",
        "salesforce", "hubspot", "zoho", "mailchimp", "google workspace", "office 365",
        "financial analysis", "forecasting", "budgeting", "accounting", "quickbooks", "sap",
    ],
}


class CVAnalyzer:
    """Parse job descriptions and extract matched skills."""

    def __init__(self):
        cfg = get_config().nlp
        self.nlp = _load_spacy_model(cfg.spacy_model, cfg.auto_download)
        self.skill_categories = SKILL_CATEGORIES

    def extract_skills_from_description(
        self, description: str
    ) -> Tuple[List[str], List[str], Dict[str, List[str]]]:
        if not description:
            logger.warning("Empty job description provided")
            return [], [], {}

        doc = self.nlp(description.lower())

        potential_skills: List[str] = []
        for chunk in doc.noun_chunks:
            potential_skills.append(chunk.text)
        for token in doc:
            if token.pos_ in ["NOUN", "PROPN"]:
                potential_skills.append(token.text)

        cleaned_skills: List[str] = []
        for skill in potential_skills:
            cleaned = re.sub(r"[^\w\s]", "", skill).strip()
            if cleaned and len(cleaned) > 1:
                cleaned_skills.append(cleaned)

        matched_skills: set = set()
        matched_requirements: List[str] = []
        matched_categories: Dict[str, List[str]] = {}

        all_skills_flat: Dict[str, str] = {}
        for category, skills in self.skill_categories.items():
            for skill in skills:
                all_skills_flat[skill] = category

        for skill in cleaned_skills:
            if skill in all_skills_flat:
                matched_skills.add(skill)
                category = all_skills_flat[skill]
                matched_categories.setdefault(category, [])
                if skill not in matched_categories[category]:
                    matched_categories[category].append(skill)
            else:
                for known_skill, category in all_skills_flat.items():
                    if known_skill in skill:
                        matched_skills.add(known_skill)
                        matched_categories.setdefault(category, [])
                        if known_skill not in matched_categories[category]:
                            matched_categories[category].append(known_skill)

        requirement_keywords = ["required", "must", "should", "need", "essential", "necessary"]
        for sent in doc.sents:
            has_modal = any(token.pos_ == "AUX" and token.dep_ == "aux" for token in sent)
            has_requirement = any(kw in sent.text.lower() for kw in requirement_keywords)
            if has_modal or has_requirement:
                matched_requirements.append(sent.text)

        if not matched_skills:
            logger.warning("No direct skill matches, trying token-level scan")
            for word in doc:
                word_text = word.text.lower()
                if word_text in all_skills_flat:
                    matched_skills.add(word_text)
                    category = all_skills_flat[word_text]
                    matched_categories.setdefault(category, [])
                    if word_text not in matched_categories[category]:
                        matched_categories[category].append(word_text)

        if not matched_skills and hasattr(doc, "user_data") and "job_title" in doc.user_data:
            job_title = doc.user_data["job_title"].lower()
            logger.warning("Falling back to generic skills for job title: %s", job_title)
            title_to_skills = {
                "developer": ["Programming Languages", "Frameworks & Libraries"],
                "engineer": ["Programming Languages", "Cloud & DevOps"],
                "data": ["Business & Analytics", "Programming Languages"],
                "analyst": ["Business & Analytics", "Databases"],
                "manager": ["Methodologies", "Soft Skills"],
                "designer": ["Tools & Platforms", "Soft Skills"],
                "marketing": ["Business & Analytics", "Soft Skills"],
                "sales": ["Business & Analytics", "Soft Skills"],
                "support": ["Soft Skills", "Tools & Platforms"],
                "admin": ["Tools & Platforms", "Business & Analytics"],
            }
            for keyword, categories in title_to_skills.items():
                if keyword in job_title:
                    for category in categories:
                        matched_categories.setdefault(category, [])
                        for skill in self.skill_categories[category][:3]:
                            if skill not in matched_categories[category]:
                                matched_categories[category].append(skill)
                                matched_skills.add(skill)

        if not matched_skills:
            logger.warning("No skills matched — using generic soft skills")
            category = "Soft Skills"
            matched_categories[category] = self.skill_categories[category][:5]
            matched_skills.update(matched_categories[category])

        logger.info(
            "Extracted %d skills across %d categories",
            len(matched_skills),
            len(matched_categories),
        )
        return list(matched_skills), matched_requirements, matched_categories

    def process_job_descriptions(
        self, jobs_df: pd.DataFrame, desc_col: str = "description", title_col: str = "title"
    ) -> pd.DataFrame:
        if jobs_df.empty:
            logger.warning("Empty DataFrame provided")
            return jobs_df

        skills_list: List[List[str]] = []
        requirements_list: List[List[str]] = []
        categories_list: List[Dict[str, List[str]]] = []

        for _, row in jobs_df.iterrows():
            description_text = str(row.get(desc_col, ""))
            job_title = str(row.get(title_col, "No Title Provided"))

            if not description_text.strip():
                logger.warning("Empty description for job: %s", job_title)
                skills_list.append([])
                requirements_list.append([])
                categories_list.append({})
                continue

            matched_skills, matched_requirements, matched_categories = (
                self.extract_skills_from_description(description_text)
            )
            skills_list.append(matched_skills)
            requirements_list.append(matched_requirements)
            categories_list.append(matched_categories)
            logger.info("Processed %s — %d skills", job_title, len(matched_skills))

        jobs_df["Extracted Skills"] = skills_list
        jobs_df["Extracted Requirements"] = requirements_list
        jobs_df["Skill Categories"] = categories_list
        return jobs_df


class CVModifier:
    """Rewrite the skills section of a .docx CV template."""

    def __init__(self, cv_template_path: str):
        self.cv_template_path = cv_template_path
        self.doc: Optional[Document] = None
        self.load_template()

    def load_template(self) -> None:
        try:
            self.doc = Document(self.cv_template_path)
            logger.info("Loaded CV template from %s", self.cv_template_path)
        except Exception as exc:
            logger.error("Error loading CV template: %s", exc)
            raise

    def find_skills_section(self):
        skill_section_keywords = [
            "skills", "technical skills", "core skills", "key skills",
            "competencies", "expertise", "qualifications", "proficiencies",
            "abilities", "capabilities", "technical competencies",
            "professional skills", "skill set", "technical expertise",
            "core competencies",
        ]

        for i, para in enumerate(self.doc.paragraphs):
            text = para.text.lower().strip()
            if text in skill_section_keywords or any(
                text.startswith(kw) for kw in skill_section_keywords
            ):
                logger.info("Found skills section at paragraph %d: '%s'", i, para.text)
                return i, para

        for i, para in enumerate(self.doc.paragraphs):
            text = para.text.lower().strip()
            if any(kw in text for kw in skill_section_keywords):
                logger.info("Found potential skills section at paragraph %d: '%s'", i, para.text)
                return i, para

        skill_related_words = [
            "proficient", "experienced", "knowledge", "familiar",
            "expert", "advanced", "intermediate", "beginner",
        ]
        for i, para in enumerate(self.doc.paragraphs):
            text = para.text.lower().strip()
            if text.startswith(("•", "-", "*")) and any(w in text for w in skill_related_words):
                if i > 0:
                    logger.info("Found skills bullet at paragraph %d", i)
                    return i - 1, self.doc.paragraphs[i - 1]

        logger.warning("Could not find skills section in CV template")
        return None, None

    def update_skills_section(self, matched_categories: Dict[str, List[str]]) -> bool:
        if not matched_categories:
            logger.warning("No matched categories provided")
            return False

        skills_idx, skills_para = self.find_skills_section()

        if skills_idx is None:
            logger.info("Creating new skills section in CV")
            skills_para = self.doc.add_paragraph()
            skills_para.style = "Heading 2"
            skills_para.text = "Skills"
            skills_idx = len(self.doc.paragraphs) - 1

        next_section_idx = None
        for i in range(skills_idx + 1, len(self.doc.paragraphs)):
            if self.doc.paragraphs[i].style.name.startswith("Heading"):
                next_section_idx = i
                break
        if next_section_idx is None:
            next_section_idx = len(self.doc.paragraphs)

        for i in reversed(range(skills_idx + 1, next_section_idx)):
            if i < len(self.doc.paragraphs):
                p = self.doc.paragraphs[i]
                p._element.getparent().remove(p._element)

        for category, skills in matched_categories.items():
            if not skills:
                continue
            category_para = self.doc.add_paragraph()
            category_para.style = "Heading 3"
            category_para.text = category
            for skill in skills:
                skill_para = self.doc.add_paragraph()
                skill_para.style = "List Bullet"
                skill_para.text = skill

        logger.info("Updated skills section in CV")
        return True

    def save_modified_cv(self, output_path: str) -> bool:
        try:
            parent = os.path.dirname(output_path)
            if parent:
                ensure_directory_exists(parent)
            self.doc.save(output_path)
            logger.info("Saved modified CV to %s", output_path)
            return True
        except Exception as exc:
            logger.error("Error saving modified CV: %s", exc)
            return False

    def process_multiple_jobs(
        self, jobs_df: pd.DataFrame, output_dir: Optional[str] = None
    ) -> List[str]:
        if jobs_df.empty:
            logger.warning("Empty DataFrame provided")
            return []

        output_dir = output_dir or str(get_config().paths.cvs_dir)
        ensure_directory_exists(output_dir)
        today_date = datetime.today().strftime("%Y-%m-%d")

        generated_cvs: List[str] = []

        for idx, row in jobs_df.iterrows():
            job_title = row.get("title", f"Job_{idx + 1}")
            company = row.get("company", "Company")
            matched_categories = row.get("Skill Categories", {})

            if not matched_categories:
                logger.warning("No skills for %s at %s", job_title, company)
                continue

            self.doc = Document(self.cv_template_path)

            if self.update_skills_section(matched_categories):
                safe_company = sanitize_filename(str(company))
                safe_title = sanitize_filename(str(job_title))
                filename = f"CV_{today_date}_{safe_company}_{safe_title}.docx"
                output_path = os.path.join(output_dir, filename)

                if self.save_modified_cv(output_path):
                    generated_cvs.append(output_path)
                    logger.info("Generated CV for %s at %s", job_title, company)
            else:
                logger.warning("Failed to update CV for %s at %s", job_title, company)

        return generated_cvs


def batch_process_jobs(
    jobs_file: str, cv_template: str, output_dir: Optional[str] = None
) -> List[str]:
    try:
        jobs_df = pd.read_excel(jobs_file)
        if jobs_df.empty:
            logger.warning("No jobs found in %s", jobs_file)
            return []

        analyzer = CVAnalyzer()
        processed_df = analyzer.process_job_descriptions(jobs_df)

        modifier = CVModifier(cv_template)
        return modifier.process_multiple_jobs(processed_df, output_dir)
    except Exception as exc:
        logger.error("Error in batch processing: %s", exc)
        return []


def main():
    jobs_file = input("Enter path to jobs Excel file: ")
    if not os.path.exists(jobs_file):
        print(f"File not found: {jobs_file}")
        return

    cv_template = input("Enter path to your CV template (.docx): ")
    if not os.path.exists(cv_template):
        print(f"File not found: {cv_template}")
        return

    output_dir = str(get_config().paths.cvs_dir)
    generated_cvs = batch_process_jobs(jobs_file, cv_template, output_dir)

    if generated_cvs:
        print(f"\nGenerated {len(generated_cvs)} tailored CVs:")
        for cv_path in generated_cvs:
            print(f"  - {cv_path}")
    else:
        print("\nFailed to generate any CVs")


if __name__ == "__main__":
    main()
